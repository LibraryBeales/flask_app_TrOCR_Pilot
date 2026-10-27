#!/usr/bin/env python3
"""
contar_errores.py

Compares every run (JSON file) of a single page against a human-reviewed
transcription (.txt) and counts the character errors (Levenshtein distance)
of each line.

It writes two tables to the output folder:
    errores_ocr_text.xlsx            -> errors of `ocr_text`
    errores_ocr_text_corrected.xlsx  -> errors of `ocr_text_corrected`

Rows = lines of the human reference; columns = runs (one JSON each).
An empty cell means that line was not detected in that run.

Usage:
    python contar_errores.py FOLDER
    python contar_errores.py FOLDER --referencia revision.txt --salida results
    python contar_errores.py FOLDER --modo-llm linea --conservar-espacios
    python contar_errores.py FOLDER --sin-s-larga --confusiones

Expected folder layout (one folder per page):
    FOLDER/
        revision.txt     <- human transcription, one line per printed line
        run_a.json       <- any file name works, as long as it ends in .json
        run_b.json

Dependencies: pandas, numpy, scipy, openpyxl  (pip install openpyxl if missing)
"""

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import linear_sum_assignment

# --- Line-matching parameters ----------------------------------------------
UMBRAL_ANCLA = 0.35      # max cost to accept a confident match (anchor)
UMBRAL_DEBIL = 0.60      # max cost for matches without positional support
PESO_LONGITUD = 0.5      # penalty when an OCR line is much longer/shorter than the reference
COSTE_NO_DETECTADA = 1.0 # cost of leaving a reference line unmatched inside a gap
SOLAPE_X_MIN = 0.5       # fraction of a box width that must fall inside the column
LARGO_MIN = 5            # shorter lines are only matched when the cost is confident


# --- Normalization ---------------------------------------------------------
def normalizar(texto, cfg, s_larga=None):
    texto = unicodedata.normalize("NFC", texto or "")
    # Convert long s to f unless disabled (s_larga overrides the global setting)
    if cfg.s_larga if s_larga is None else s_larga:
        texto = texto.replace("ſ", "f")
    if not cfg.conservar_espacios:
        texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


# --- Distances -------------------------------------------------------------
def levenshtein(a, b):
    if len(a) < len(b):
        a, b = b, a
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def alinear(ref, hyp):
    """Character-by-character alignment: list of (ref_char, hyp_char) pairs.
    None means an insertion (missing in ref) or a deletion (missing in hyp)."""
    n, m = len(ref), len(hyp)
    D = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n + 1):
        D[i][0] = i
    for j in range(m + 1):
        D[0][j] = j
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            D[i][j] = min(D[i - 1][j] + 1, D[i][j - 1] + 1,
                          D[i - 1][j - 1] + (ref[i - 1] != hyp[j - 1]))
    pares, i, j = [], n, m
    while i > 0 or j > 0:
        if i > 0 and j > 0 and D[i][j] == D[i - 1][j - 1] + (ref[i - 1] != hyp[j - 1]):
            pares.append((ref[i - 1], hyp[j - 1]))
            i, j = i - 1, j - 1
        elif i > 0 and D[i][j] == D[i - 1][j] + 1:
            pares.append((ref[i - 1], None))
            i -= 1
        else:
            pares.append((None, hyp[j - 1]))
            j -= 1
    return pares[::-1]


# Row labels of the long-s confusion report (kept in Spanish: they are output data)
CATEGORIAS_S = [
    "ſ leída como ſ (correcto)",
    "ſ leída como f",
    "ſ leída como s",
    "ſ leída como otro carácter",
    "ſ no leída (borrada)",
    "f real leída como f (correcto)",
    "f real leída como ſ",
    "f real leída como otro carácter",
]


def contar_confusiones_s(ref_n, hyp_n):
    """Counts how the long s (ſ) and the real f characters of a line were read."""
    c = dict.fromkeys(CATEGORIAS_S, 0)
    for r, h in alinear(ref_n, hyp_n):
        if r == "ſ":
            if h == "ſ":
                c[CATEGORIAS_S[0]] += 1
            elif h == "f":
                c[CATEGORIAS_S[1]] += 1
            elif h == "s":
                c[CATEGORIAS_S[2]] += 1
            elif h is None:
                c[CATEGORIAS_S[4]] += 1
            else:
                c[CATEGORIAS_S[3]] += 1
        elif r == "f":
            if h == "f":
                c[CATEGORIAS_S[5]] += 1
            elif h == "ſ":
                c[CATEGORIAS_S[6]] += 1
            else:
                c[CATEGORIAS_S[7]] += 1
    return c


def distancia_ajuste(patron, texto):
    """Minimum distance between `patron` and ANY fragment of `texto`.
    Lets us recognize a reference line even if the OCR appended garbage to it."""
    if not patron:
        return 0
    prev = [0] * (len(texto) + 1)
    for i, cp in enumerate(patron, 1):
        cur = [i]
        for j, ct in enumerate(texto, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (cp != ct)))
        prev = cur
    return min(prev)


def coste_emparejamiento(ref, hyp):
    """Matching cost between a reference line and an OCR line (lower is better)."""
    if not ref:
        return 9.0
    ajuste = distancia_ajuste(ref, hyp) / len(ref)
    dif_long = abs(len(hyp) - len(ref)) / max(len(hyp), len(ref), 1)
    return ajuste + PESO_LONGITUD * dif_long


def errores_por_linea_continuo(ref_lineas, hyp_lineas):
    """Aligns the continuous text (no line breaks, no end-of-line hyphens) and
    attributes every error to the reference line where it occurs.
    ref_lineas: list of (line_id, text); hyp_lineas: list of texts."""
    def aplanar(pares):
        chars = []
        for lid, t in pares:
            if t.endswith("-"):
                chars += [(c, lid) for c in t[:-1]]
            else:
                chars += [(c, lid) for c in t] + [(" ", lid)]
        return chars

    ref = aplanar(ref_lineas)
    hyp = [c for c, _ in aplanar([(None, t) for t in hyp_lineas])]
    n, m = len(ref), len(hyp)
    D = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n + 1):
        D[i][0] = i
    for j in range(m + 1):
        D[0][j] = j
    for i in range(1, n + 1):
        ri = ref[i - 1][0]
        fila, arriba = D[i], D[i - 1]
        for j in range(1, m + 1):
            fila[j] = min(arriba[j] + 1, fila[j - 1] + 1,
                          arriba[j - 1] + (ri != hyp[j - 1]))

    errores = {lid: 0 for lid, _ in ref_lineas}
    i, j = n, m
    while i > 0 or j > 0:
        if i > 0 and j > 0 and D[i][j] == D[i - 1][j - 1] + (ref[i - 1][0] != hyp[j - 1]):
            if ref[i - 1][0] != hyp[j - 1]:
                errores[ref[i - 1][1]] += 1
            i, j = i - 1, j - 1
        elif i > 0 and D[i][j] == D[i - 1][j] + 1:
            errores[ref[i - 1][1]] += 1          # deletion
            i -= 1
        else:
            lid = ref[i - 1][1] if i > 0 else ref[0][1]
            errores[lid] += 1                    # insertion
            j -= 1
    return errores


# --- File reading ----------------------------------------------------------
def leer_referencia(ruta):
    lineas = ruta.read_text(encoding="utf-8-sig").splitlines()
    return [l.rstrip() for l in lineas if l.strip()]


def leer_ejecucion(ruta):
    with open(ruta, encoding="utf-8") as f:
        datos = json.load(f)
    segmentos = []
    for k, s in enumerate(datos.get("line_segments", [])):
        bbox = s.get("bbox") or []
        segmentos.append({
            "orden": s.get("line_index", k),
            "ocr": s.get("ocr_text", "") or "",
            "llm": s.get("ocr_text_corrected", "") or "",
            "bbox": bbox if len(bbox) == 4 else None,
        })
    return datos.get("timestamp", 0), segmentos


# --- Line matching ---------------------------------------------------------
def _y(seg):
    b = seg["bbox"]
    return (b[1] + b[3]) / 2 if b else seg["orden"]


def _en_columna(seg, x1, x2):
    b = seg["bbox"]
    if not b or x1 is None:
        return True
    ancho = max(b[2] - b[0], 1)
    solape = max(0, min(b[2], x2) - max(b[0], x1))
    return solape / ancho >= SOLAPE_X_MIN


def emparejar(ref, segs, cfg):
    """Returns {reference_line_index: segment_index}."""
    ref_n = [normalizar(r, cfg) for r in ref]
    hyp_n = [normalizar(s["ocr"], cfg) for s in segs]
    if not segs:
        return {}
    C = np.array([[coste_emparejamiento(r, h) for h in hyp_n] for r in ref_n])
    # A very short line ("V.", "El") resembles any stamp or mark, so we only
    # accept it when the similarity is clear.
    for i, r in enumerate(ref_n):
        if len(r) < LARGO_MIN:
            C[i, C[i] > UMBRAL_ANCLA] = 9.0

    # 1) Anchors: optimal global assignment, keeping only the confident pairs.
    filas, cols = linear_sum_assignment(C)
    hungaro = dict(zip(filas, cols))
    pares = {i: j for i, j in hungaro.items() if C[i, j] <= UMBRAL_ANCLA}
    usados = set(pares.values())

    # 2) Gaps: reference lines left unmatched between two anchors are searched
    #    among the boxes that sit physically between those anchors.
    i = 0
    while i < len(ref):
        if i in pares:
            i += 1
            continue
        fin = i
        while fin < len(ref) and fin not in pares:
            fin += 1
        hueco = list(range(i, fin))
        ant = segs[pares[i - 1]] if i - 1 in pares else None
        sig = segs[pares[fin]] if fin in pares else None
        cajas = [s["bbox"] for s in (ant, sig) if s and s["bbox"]]
        x1 = min(b[0] for b in cajas) if cajas else None
        x2 = max(b[2] for b in cajas) if cajas else None
        cand = [j for j, s in enumerate(segs)
                if j not in usados
                and (ant is None or _y(s) > _y(ant))
                and (sig is None or _y(s) < _y(sig))
                and _en_columna(s, x1, x2)]
        cand.sort(key=lambda j: _y(segs[j]))

        # Monotonic DP: assign the gap's lines to candidates in vertical order.
        k, m = len(hueco), len(cand)
        INF = float("inf")
        dp = [[INF] * (m + 1) for _ in range(k + 1)]
        paso = [[None] * (m + 1) for _ in range(k + 1)]
        dp[0] = [0.0] * (m + 1)
        for a in range(1, k + 1):
            dp[a][0] = dp[a - 1][0] + COSTE_NO_DETECTADA
            paso[a][0] = "sin"
        for a in range(1, k + 1):
            for b in range(1, m + 1):
                opciones = [
                    (dp[a][b - 1], "salta"),                                  # skip a candidate box
                    (dp[a - 1][b] + COSTE_NO_DETECTADA, "sin"),               # line not detected
                    (dp[a - 1][b - 1] + C[hueco[a - 1], cand[b - 1]], "par"), # match line to box
                ]
                dp[a][b], paso[a][b] = min(opciones)
        a, b = k, m
        while a > 0:
            p = paso[a][b] if b > 0 else "sin"
            if p == "par":
                pares[hueco[a - 1]] = cand[b - 1]
                usados.add(cand[b - 1])
                a, b = a - 1, b - 1
            elif p == "salta":
                b -= 1
            else:
                a -= 1
        i = fin

    # 3) Leftovers: acceptable Hungarian matches even without positional support.
    for i, j in hungaro.items():
        if i not in pares and j not in usados and C[i, j] <= UMBRAL_DEBIL:
            pares[i] = j
            usados.add(j)
    return pares


# --- Main program ----------------------------------------------------------
def procesar(carpeta, cfg):
    ref_ruta = Path(cfg.referencia) if cfg.referencia else None
    if ref_ruta is None:
        txts = sorted(carpeta.glob("*.txt"))
        if len(txts) != 1:
            sys.exit(f"Expected exactly one .txt in {carpeta}, found {len(txts)}. "
                     "Choose one with --referencia.")
        ref_ruta = txts[0]
    ref = leer_referencia(ref_ruta)
    ref_n = [normalizar(r, cfg) for r in ref]

    ejecuciones = []
    for ruta in sorted(carpeta.glob("*.json")):
        try:
            ts, segs = leer_ejecucion(ruta)
        except (json.JSONDecodeError, UnicodeDecodeError) as e:
            print(f"Skipped {ruta.name}: {e}")
            continue
        if segs:
            ejecuciones.append((ts, ruta.stem, segs))
    if not ejecuciones:
        sys.exit("No JSON file with 'line_segments' found in the folder.")
    # Columns are ordered by the JSON timestamp, then by file name
    ejecuciones.sort(key=lambda e: (e[0], e[1]))

    etiquetas = [f"{i + 1:02d} {r[:40]}" for i, r in enumerate(ref)]
    tabla_ocr = pd.DataFrame(index=etiquetas, dtype="Int64")
    tabla_llm = pd.DataFrame(index=etiquetas, dtype="Int64")
    total_ref = sum(len(r) for r in ref_n)
    conf = {"ocr_text": {}, "ocr_text_corrected": {}}
    if cfg.confusiones:
        ref_s = [normalizar(r, cfg, s_larga=False) for r in ref]
        if not any("ſ" in r for r in ref_s):
            print("Warning: the reference contains no 'ſ'; the confusion report will be empty.")

    print(f"Reference: {ref_ruta.name} ({len(ref)} lines)")
    for _, nombre, segs in ejecuciones:
        pares = emparejar(ref, segs, cfg)
        col_ocr, col_llm = [pd.NA] * len(ref), [pd.NA] * len(ref)

        for i, j in pares.items():
            col_ocr[i] = levenshtein(ref_n[i], normalizar(segs[j]["ocr"], cfg))

        orden = sorted(pares)
        if cfg.modo_llm == "linea":
            # Line by line, exactly as stored in the JSON
            for i in orden:
                col_llm[i] = levenshtein(ref_n[i], normalizar(segs[pares[i]]["llm"], cfg))
        elif orden:
            # Continuous text: tolerates the LLM moving text between lines
            errores = errores_por_linea_continuo(
                [(i, ref_n[i]) for i in orden],
                [normalizar(segs[pares[i]]["llm"], cfg) for i in orden])
            for i, e in errores.items():
                col_llm[i] = e

        if cfg.confusiones:
            for campo, clave in [("ocr", "ocr_text"), ("llm", "ocr_text_corrected")]:
                total = dict.fromkeys(CATEGORIAS_S, 0)
                if campo == "ocr" or cfg.modo_llm == "linea":
                    for i, j in pares.items():
                        c = contar_confusiones_s(
                            ref_s[i], normalizar(segs[j][campo], cfg, s_larga=False))
                        for k, v in c.items():
                            total[k] += v
                else:
                    # Continuous text: align the whole block of matched lines
                    ids = sorted(pares)
                    ref_txt = " ".join(ref_s[i].rstrip("-") for i in ids)
                    hyp_txt = " ".join(normalizar(segs[pares[i]][campo], cfg, s_larga=False).rstrip("-")
                                       for i in ids)
                    total = contar_confusiones_s(ref_txt, hyp_txt)
                conf[clave][nombre] = total

        tabla_ocr[nombre] = pd.array(col_ocr, dtype="Int64")
        tabla_llm[nombre] = pd.array(col_llm, dtype="Int64")

        detectadas = sum(len(ref_n[i]) for i in pares)
        e_ocr, e_llm = tabla_ocr[nombre].sum(), tabla_llm[nombre].sum()
        print(f"  {nombre}: {len(pares)}/{len(ref)} lines | "
              f"OCR {e_ocr} err ({e_ocr / max(detectadas, 1):.1%}) | "
              f"LLM {e_llm} err ({e_llm / max(detectadas, 1):.1%})")

    salida = Path(cfg.salida) if cfg.salida else carpeta
    salida.mkdir(parents=True, exist_ok=True)
    for tabla, nombre_arch in [(tabla_ocr, "errores_ocr_text.xlsx"),
                               (tabla_llm, "errores_ocr_text_corrected.xlsx")]:
        ruta = salida / nombre_arch
        guardar_excel(tabla, ruta)
        print(f"Saved {ruta}")
    if total_ref == 0:
        print("Warning: the reference is empty.")
    if cfg.confusiones:
        ruta = salida / "confusiones_s_larga.xlsx"
        with pd.ExcelWriter(ruta, engine="openpyxl") as w:
            for hoja, datos in conf.items():
                df = pd.DataFrame(datos, index=CATEGORIAS_S)
                df.to_excel(w, index_label="Categoría", sheet_name=hoja)
        formatear_hojas(ruta)
        print(f"Saved {ruta}")


def guardar_excel(tabla, ruta):
    from openpyxl.styles import Alignment, Font
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        tabla.to_excel(w, index_label="Línea", sheet_name="Errores")
        hoja = w.sheets["Errores"]
        for fila in hoja.iter_rows():
            for celda in fila:
                celda.font = Font(name="Arial", size=10, bold=celda.row == 1)
                if celda.column > 1:
                    celda.alignment = Alignment(horizontal="center", wrap_text=celda.row == 1)
        hoja.column_dimensions["A"].width = 48
        for col in range(2, hoja.max_column + 1):
            hoja.column_dimensions[hoja.cell(1, col).column_letter].width = 22
        hoja.freeze_panes = "B2"


def formatear_hojas(ruta):
    """Applies the same plain formatting to every sheet of an existing workbook."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    wb = load_workbook(ruta)
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                celda.font = Font(name="Arial", size=10, bold=celda.row == 1)
                if celda.column > 1:
                    celda.alignment = Alignment(horizontal="center", wrap_text=celda.row == 1)
        hoja.column_dimensions["A"].width = 36
        for col in range(2, hoja.max_column + 1):
            hoja.column_dimensions[hoja.cell(1, col).column_letter].width = 22
        hoja.freeze_panes = "B2"
    wb.save(ruta)


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("carpeta", help="Folder with the JSON runs of one page and its reviewed .txt")
    p.add_argument("--referencia", help="Path of the reviewed .txt (needed if the folder has several)")
    p.add_argument("--salida", help="Output folder (default: the input folder)")
    p.add_argument("--modo-llm", choices=["continuo", "linea"], default="continuo",
                   help="continuo: tolerates the LLM moving text between lines (default); "
                        "linea: compares each line exactly as stored")
    p.add_argument("--conservar-espacios", action="store_true",
                   help="Count double spaces as errors")
    p.add_argument("--sin-s-larga", dest="s_larga", action="store_false",
                   help="Do not convert 'ſ' to 'f' before comparing")
    p.add_argument("--confusiones", action="store_true",
                   help="Write confusiones_s_larga.xlsx: how each ſ was read (as ſ, f, s...). "
                        "Use together with --sin-s-larga and a reference that contains 'ſ'")
    cfg = p.parse_args()
    carpeta = Path(cfg.carpeta)
    if not carpeta.is_dir():
        sys.exit(f"Folder not found: {carpeta}")
    procesar(carpeta, cfg)


if __name__ == "__main__":
    main()
